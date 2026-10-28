import base64
import copy
import json
import math
import os
import tempfile
from pathlib import Path

import openpyxl
import pytest

_tmp = tempfile.mkdtemp()
os.environ["DATABASE_URL"] = f"sqlite:///{Path(_tmp) / 'test.db'}"
os.environ["DASHBOARD_USERS"] = "agent:secret,superviseur:pw2"
os.environ["KOBO_WEBHOOK_TOKEN"] = "tok123"
os.environ["SECRET_KEY"] = "test-key"

from fastapi.testclient import TestClient  # noqa: E402

from app.main import EXEMPLE_MENAGE, app  # noqa: E402
from app.schema import FEATURE_NAMES, FEATURES  # noqa: E402
from app.security import csrf_token  # noqa: E402
from kobo.build_xlsform import GROUPS  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
GROUPE = {name: groupe for groupe, _, names in GROUPS for name in names}


def kobo(reponses: dict) -> dict:
    """Réponses -> questions préfixées par leur groupe, valeurs en texte, comme un REST Service Kobo."""
    return {f"{GROUPE[k]}/{k}": str(v) for k, v in reponses.items()}


# Soumission telle qu'envoyée par un REST Service KoboToolbox (groupes préfixés, valeurs en texte)
KOBO_PAYLOAD = {
    "_id": 4242,
    "_uuid": "b1f0c6a2-0000-4c3b-9d1e-123456789abc",
    "_submission_time": "2026-10-20T09:15:02",
    "_xform_id_string": "ciblage_menages_eligibilite",
    "_geolocation": [-1.2, 29.4],
    "_attachments": [],
    "formhub/uuid": "abc",
    "consentement": "oui",
    "identification/enumerateur": "Grace",
    "identification/village": "Kibirizi centre",
    "identification/nom_chef_menage": "Amani Kahindo",
    "identification/telephone": "+243 990 000 000",
    "identification/gps": "-1.2 29.4 1500 5",
    **kobo(EXEMPLE_MENAGE),
    "meta/instanceID": "uuid:b1f0c6a2-0000-4c3b-9d1e-123456789abc",
}

NON_VULNERABLE = {
    **EXEMPLE_MENAGE,
    "localite": "Rutshuru", "statut_menage": "resident", "duree_deplacement_mois": None, "nombre_deplacements_12m": None,
    "sexe_chef_menage": "homme", "age_chef_menage": 40, "etat_civil_chef": "marie", "instruction_chef": "secondaire",
    "taille_menage": 4, "enfants_moins_5": 1, "enfants_5_17": 1, "enfants_non_scolarises": 0, "femme_enceinte_allaitante": "non",
    "type_logement": "maison_correcte", "statut_occupation": "proprietaire", "nombre_pieces": 3, "articles_essentiels": 5,
    "possede_betail": "oui", "acces_terre_cultivable": "oui", "toit_etanche": "oui", "source_eclairage": "solaire",
    "possede_telephone": "oui", "source_eau": "robinet", "temps_collecte_eau_min": 10, "type_latrine": "privee",
    "savon_disponible": "oui", "enfant_malade_2_semaines": "non", "acces_soins": "oui", "vaccination_enfants_a_jour": "oui",
    "jours_cereales": 7, "jours_legumineuses": 5, "jours_lait": 2, "jours_viande_poisson": 3, "jours_legumes": 6,
    "jours_fruits": 3, "jours_huile": 7, "jours_sucre": 5, "stock_nourriture_jours": 30, "repas_par_jour": 3,
    "jours_sans_nourriture_7j": 0, "rcsi_aliments_moins_chers": 0, "rcsi_emprunt_nourriture": 0, "rcsi_reduction_portions": 0,
    "rcsi_adultes_pour_enfants": 0, "rcsi_moins_de_repas": 0, "lcs_emprunt_argent": "non", "lcs_vente_biens_menagers": "non",
    "lcs_reduction_depenses_sante": "non", "lcs_retrait_enfants_ecole": "non", "source_revenu": "petit_commerce",
    "membres_avec_revenu": 2, "perte_revenu_6_mois": "non", "revenu_30j_usd": 220, "part_depenses_alimentation_pct": 40,
    "dette_usd": 0, "epargne_disponible": "oui", "incident_securite_3_mois": "non", "biens_perdus_crise": "non",
    "documents_identite": "oui", "restriction_mouvements": "non",
}
NON_VULNERABLE = {k: v for k, v in NON_VULNERABLE.items() if v is not None}


def auth(user="agent", pw="secret"):
    return {"Authorization": "Basic " + base64.b64encode(f"{user}:{pw}".encode()).decode()}


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


def test_health(client):
    assert client.get("/health").json()["status"] == "ok"


def test_webhook_requires_token(client):
    assert client.post("/api/kobo/webhook", json=KOBO_PAYLOAD).status_code == 401
    assert client.post("/api/kobo/webhook?token=faux", json=KOBO_PAYLOAD).status_code == 401


def test_webhook_scores_and_is_idempotent(client):
    r = client.post("/api/kobo/webhook?token=tok123", json=KOBO_PAYLOAD)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["cree"] is True and body["eligible"] is True
    again = client.post("/api/kobo/webhook", json=KOBO_PAYLOAD, headers={"X-Webhook-Token": "tok123"})
    assert again.json() == {**body, "cree": False}
    assert len(client.get("/api/beneficiaires", headers=auth()).json()) == 1


def test_webhook_rejects_invalid_answers(client):
    bad = {**KOBO_PAYLOAD, "_id": 1, "profil/localite": "Goma", "profil/age_chef_menage": "abc"}
    r = client.post("/api/kobo/webhook?token=tok123", json=bad)
    assert r.status_code == 422
    assert "localite" in r.json()["detail"] and "age_chef_menage" in r.json()["detail"]


def test_questions_conditionnelles_et_coherence(client):
    # déplacé sans durée de déplacement : refusé
    sans_duree = {k: v for k, v in KOBO_PAYLOAD.items() if k != "profil/duree_deplacement_mois"}
    assert "duree_deplacement_mois" in client.post("/api/predict", json=sans_duree).json()["detail"]
    # résident : durée et nombre de déplacements non demandés ; assistance reçue : le type devient obligatoire
    assert client.post("/api/predict", json=NON_VULNERABLE).status_code == 200
    resident = {**NON_VULNERABLE, "assistance_3_derniers_mois": "oui"}
    assert "type_assistance_recue" in client.post("/api/predict", json=resident).json()["detail"]
    assert client.post("/api/predict", json={**resident, "type_assistance_recue": "abri"}).status_code == 200
    # plus d'enfants et de personnes âgées que de membres, plus d'enfants non scolarisés que d'enfants : refusé
    trop = client.post("/api/predict", json={**NON_VULNERABLE, "enfants_5_17": 4}).json()["detail"]
    assert "personnes_60_plus" in trop
    assert "enfants_non_scolarises" in client.post("/api/predict", json={**NON_VULNERABLE, "enfants_non_scolarises": 2}).json()["detail"]


def test_webhook_ignores_refused_consent(client):
    r = client.post("/api/kobo/webhook?token=tok123", json={"_id": 9, "consentement": "non"})
    assert r.status_code == 200 and r.json()["ignore"] is True


def test_predict_explication_additive(client):
    r = client.post("/api/predict", json=NON_VULNERABLE).json()
    assert r["eligible"] is False and r["score"] < r["threshold"]
    # base + somme des contributions = log-odds, et probabilité = sigmoïde(log-odds)
    assert abs(r["base_value"] + sum(e["contribution"] for e in r["explanation"]) - r["log_odds"]) < 1e-2
    assert abs(1 / (1 + math.exp(-r["log_odds"])) - r["score"]) < 1e-3
    assert {e["feature"] for e in r["explanation"]} == set(FEATURE_NAMES)
    assert abs(sum(t["contribution"] for t in r["themes"]) - sum(e["contribution"] for e in r["explanation"])) < 1e-2
    assert r["indice_priorite"] == pytest.approx(r["log_odds"], abs=1e-3)


def test_pas_de_grille_dans_le_modele(client):
    meta = client.get("/api/modele").json()
    assert "seuil_grille" not in meta and "score_grille_excel" not in meta["variables"]
    assert len(meta["variables"]) > 70
    # objectif du projet : au moins 95 % sur toutes les métriques du jeu de test
    for m in ["accuracy", "f1", "rappel", "precision"]:
        assert meta["test"][m] >= 0.95, m


def test_type_assistance_ideale(client):
    from app.schema import ASSISTANCES
    r = client.post("/api/predict", json=KOBO_PAYLOAD).json()
    a = r["assistance"]
    assert a["code"] in ASSISTANCES and a["libelle"] == ASSISTANCES[a["code"]]
    assert len(a["alternatives"]) == 2 and a["probabilite"] >= max(x["probabilite"] for x in a["alternatives"])
    assert a["probabilite"] + sum(x["probabilite"] for x in a["alternatives"]) <= 1.0001
    assert a["facteurs"] and all(f["contribution"] > 0 for f in a["facteurs"])
    # ménage dont le besoin principal est clairement l'éducation : 3 enfants sur 3 hors de l'école, retirés et au travail
    scolaire = {**NON_VULNERABLE, "enfants_5_17": 3, "taille_menage": 6, "enfants_non_scolarises": 3,
                "lcs_retrait_enfants_ecole": "oui", "lcs_travail_enfants": "oui"}
    assert client.post("/api/predict", json=scolaire).json()["assistance"]["code"] == "education"
    meta = client.get("/api/modele").json()["assistance"]
    assert meta["classes"] == sorted(ASSISTANCES)
    assert meta["test"]["accuracy"] >= 0.95 and meta["test"]["f1_macro"] >= 0.95


def test_migration_ajoute_les_nouvelles_colonnes(tmp_path):
    from sqlalchemy import create_engine, inspect, text
    from app.db import ajouter_colonnes_manquantes
    engine = create_engine(f"sqlite:///{tmp_path / 'ancienne.db'}")
    with engine.begin() as con:  # tables d'une version précédente, sans les colonnes « assistance »
        con.execute(text("CREATE TABLE beneficiaires (id INTEGER PRIMARY KEY)"))
        con.execute(text("CREATE TABLE selections (id INTEGER PRIMARY KEY)"))
    ajouter_colonnes_manquantes(engine)
    ajouter_colonnes_manquantes(engine)  # sans effet la deuxième fois
    for table in ("beneficiaires", "selections"):
        assert "assistance" in {c["name"] for c in inspect(engine).get_columns(table)}


def test_human_override_and_audit_trail(client):
    ben_id = client.get("/api/beneficiaires", headers=auth()).json()[0]["id"]
    r = client.post(f"/api/beneficiaires/{ben_id}/revision", json={"approuve": False}, headers=auth())
    assert r.status_code == 422
    r = client.post(f"/api/beneficiaires/{ben_id}/revision", headers=auth("superviseur", "pw2"),
                    json={"approuve": False, "commentaire": "Ménage déjà assisté par un autre partenaire"})
    d = r.json()
    assert d["decision_modele"] == "eligible" and d["decision_finale"] == "non_eligible"
    assert d["statut_revue"] == "modifiee" and d["reviewed_by"] == "superviseur"
    assert d["revues"][0]["commentaire"].startswith("Ménage déjà")


def test_kobo_edit_resets_review(client):
    edited = copy.deepcopy(KOBO_PAYLOAD)
    edited["alimentation/repas_par_jour"] = "3"
    client.post("/api/kobo/webhook?token=tok123", json=edited)
    d = client.get("/api/beneficiaires", headers=auth()).json()[0]
    assert d["statut_revue"] == "en_attente"


def test_dashboard_pages_and_csrf(client):
    assert client.get("/dashboard").status_code == 401
    assert client.get("/dashboard", headers=auth("agent", "mauvais")).status_code == 401
    page = client.get("/dashboard", headers=auth())
    assert page.status_code == 200 and "Amani Kahindo" in page.text
    assert "Assistance idéale" not in page.text and 'href="/docs"' not in page.text  # liste sans assistance, pas d'onglet API
    ben_id = client.get("/api/beneficiaires", headers=auth()).json()[0]["id"]
    detail = client.get(f"/dashboard/beneficiaires/{ben_id}", headers=auth())
    assert detail.status_code == 200 and "Pourquoi le modèle" in detail.text and "Abri de fortune" in detail.text
    assert "Par thème du questionnaire" in detail.text and "Type d'assistance idéale" in detail.text
    modele = client.get("/dashboard/modele", headers=auth())
    assert modele.status_code == 200 and "notebooks/ciblage_avec_ml.ipynb" in modele.text
    assert "grille" not in modele.text.lower()
    url = f"/dashboard/beneficiaires/{ben_id}/revision"
    assert client.post(url, data={"decision": "approuve"}, headers=auth()).status_code == 403
    r = client.post(url, data={"decision": "approuve", "csrf": csrf_token("agent")}, headers=auth(),
                    follow_redirects=False)
    assert r.status_code == 303
    csv_text = client.get("/dashboard/export.csv", headers=auth()).text
    assert "Amani Kahindo" in csv_text and "confirmee" in csv_text


def test_xlsform_matches_schema(tmp_path, monkeypatch):
    from kobo import build_xlsform
    monkeypatch.setattr(build_xlsform, "OUT", tmp_path / "formulaire.xlsx")  # ne touche pas au fichier versionné
    wb = openpyxl.load_workbook(build_xlsform.build())
    survey = list(wb["survey"].iter_rows(values_only=True))
    header = survey[0]
    rows = {r[header.index("name")]: r for r in survey[1:]}
    choices = {}
    for list_name, name, _ in list(wb["choices"].iter_rows(values_only=True))[1:]:
        choices.setdefault(list_name, set()).add(name)
    for f in FEATURES:
        assert f.name in rows, f"question {f.name} absente du formulaire"
        qtype = rows[f.name][header.index("type")]
        if f.kind == "categorical":
            assert choices[qtype.split()[1]] == set(f.choices), f.name
        if f.relevant:
            assert f"${{{f.relevant[0]}}}" in rows[f.name][header.index("relevant")], f.name
    assert "${taille_menage}" in rows["personnes_60_plus"][header.index("constraint")]


def test_metadata_versionnee_coherente():
    meta = json.loads((ROOT / "model" / "metadata.json").read_text(encoding="utf-8"))
    assert meta["variables"] == FEATURE_NAMES
