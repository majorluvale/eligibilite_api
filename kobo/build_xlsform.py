"""Génère le formulaire KoboToolbox (XLSForm) à partir de app/schema.py.

Usage :  python kobo/build_xlsform.py
Sortie : kobo/formulaire_ciblage_menages.xlsx  (à importer dans KoboToolbox)

Le formulaire est généré, pas écrit à la main : les noms des questions et
les codes des choix sont donc toujours ceux que l'API attend.
"""
from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app.schema import COHERENCE, FEATURES, FEATURES_BY_KOBO  # noqa: E402

OUT = ROOT / "kobo" / "formulaire_ciblage_menages.xlsx"
SURVEY_COLS = ["type", "name", "label", "hint", "required", "constraint", "constraint_message",
               "relevant", "appearance"]

# Groupes du formulaire : un groupe par thème de app/schema.py, dans l'ordre du schéma
NOMS_GROUPES = {"Profil": "profil", "Composition": "composition", "Logement et biens": "logement",
                "Eau et hygiène": "eau_hygiene", "Santé et nutrition": "sante", "Alimentation": "alimentation",
                "Stratégies d'adaptation": "adaptation", "Moyens d'existence": "moyens_existence",
                "Protection": "protection", "Assistance": "assistance"}
GROUPS = []
for _f in FEATURES:
    if not GROUPS or GROUPS[-1][1] != _f.theme:
        GROUPS.append((NOMS_GROUPES[_f.theme], _f.theme, []))
    GROUPS[-1][2].append(_f.name)
# Questions numériques présentées comme un choix (libellés plus clairs pour l'enquêteur)
AS_CHOICES = {"repas_par_jour": {"1": "1 repas ou moins", "2": "2 repas", "3": "3 repas ou plus"}}
# Contraintes entre questions, vérifiées par Kobo dès la saisie (l'API les vérifie aussi)
CONTRAINTES = {question: (contrainte, message[0].upper() + message[1:]) for question, contrainte, _, message in COHERENCE}
CONSENT = "${consentement} = 'oui'"


def row(**kw) -> dict:
    return kw


def build() -> Path:
    survey, choices = [], []
    shared_lists = {}

    def add_list(list_name: str, items: dict[str, str]):
        if list_name in shared_lists:
            return
        shared_lists[list_name] = items
        choices.extend({"list_name": list_name, "name": k, "label": v} for k, v in items.items())

    add_list("oui_non", {"oui": "Oui", "non": "Non"})

    survey += [
        row(type="start", name="start"),
        row(type="end", name="end"),
        row(type="today", name="today"),
        row(type="deviceid", name="deviceid"),
        row(type="note", name="intro",
            label="**Enquête de ciblage des ménages**\nCe formulaire sert à évaluer l'éligibilité du ménage "
                  "à l'assistance humanitaire. Les réponses sont analysées automatiquement puis revues par l'équipe."),
        row(type="select_one oui_non", name="consentement", required="yes",
            label="Le chef de ménage (ou son représentant) accepte-t-il de répondre à cette enquête ?",
            hint="Lire la déclaration de consentement avant de poser la question."),
        row(type="note", name="fin_sans_consentement", relevant="${consentement} = 'non'",
            label="Le ménage n'a pas donné son consentement. Remerciez la personne et terminez l'entretien."),

        row(type="begin_group", name="identification", label="Identification", relevant=CONSENT),
        row(type="text", name="enumerateur", label="Nom de l'enquêteur", required="yes"),
        row(type="text", name="village", label="Village / quartier", required="yes"),
        row(type="text", name="nom_chef_menage", label="Nom complet du chef de ménage", required="yes"),
        row(type="text", name="telephone", label="Numéro de téléphone du ménage (si disponible)",
            constraint="regex(., '^[+0-9 ]{8,16}$')",
            constraint_message="Numéro invalide : chiffres, espaces ou + uniquement (8 à 16 caractères)",
            appearance="numbers"),
        row(type="geopoint", name="gps", label="Position GPS du ménage"),
        row(type="end_group", name=""),
    ]

    for group, label, names in GROUPS:
        survey.append(row(type="begin_group", name=group, label=label, relevant=CONSENT))
        for name in names:
            f = FEATURES_BY_KOBO[name]
            relevant = ""
            if f.relevant:  # question conditionnelle (logique de saut)
                question, valeurs = f.relevant
                relevant = " or ".join(f"${{{question}}} = '{v}'" for v in sorted(valeurs))
            if name in AS_CHOICES:
                add_list(name, AS_CHOICES[name])
                survey.append(row(type=f"select_one {name}", name=name, label=f.label, hint=f.hint, required="yes",
                                  relevant=relevant))
            elif f.kind == "integer":
                contrainte, message = f". >= {f.min} and . <= {f.max}", f"La valeur doit être comprise entre {f.min} et {f.max}"
                if name in CONTRAINTES:
                    contrainte, message = f"{contrainte} and {CONTRAINTES[name][0]}", f"{message}. {CONTRAINTES[name][1]}"
                survey.append(row(type="integer", name=name, label=f.label, hint=f.hint, required="yes",
                                  constraint=contrainte, constraint_message=message, relevant=relevant))
            else:
                list_name = "oui_non" if f.choices == {"oui": "Oui", "non": "Non"} else name
                add_list(list_name, f.choices)
                survey.append(row(type=f"select_one {list_name}", name=name, label=f.label, hint=f.hint,
                                  required="yes", relevant=relevant))
        survey.append(row(type="end_group", name=""))

    survey.append(row(type="note", name="fin", relevant=CONSENT,
                      label="Merci. L'éligibilité sera calculée à la réception de la soumission et "
                            "revue par l'équipe de ciblage."))

    wb = Workbook()
    write_sheet(wb.active, "survey", SURVEY_COLS, survey)
    write_sheet(wb.create_sheet(), "choices", ["list_name", "name", "label"], choices)
    write_sheet(wb.create_sheet(), "settings", ["form_title", "form_id", "version", "default_language"], [{
        "form_title": "Ciblage des ménages - éligibilité à l'assistance",
        "form_id": "ciblage_menages_eligibilite",
        "version": datetime.now().strftime("%Y%m%d%H"),
        "default_language": "French (fr)",
    }])
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    return OUT


def write_sheet(ws, title: str, cols: list[str], rows: list[dict]):
    ws.title = title
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for r in rows:
        ws.append([r.get(c, "") for c in cols])
    for i, c in enumerate(cols, 1):
        width = max(len(str(r.get(c, ""))) for r in rows + [{c: c}])
        ws.column_dimensions[get_column_letter(i)].width = min(max(12, width + 2), 70)
    ws.freeze_panes = "A2"


if __name__ == "__main__":
    print(f"Formulaire écrit : {build()}")
